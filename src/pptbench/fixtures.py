"""Project-authored synthetic fixture corpus and its frozen task definitions."""

from __future__ import annotations

import hashlib
import json
import os
import re
import stat
import tempfile
import zipfile
from importlib.resources import files
from io import BytesIO
from pathlib import Path
from typing import Any, TypedDict

from .util import sha256_file

MANIFEST_RESOURCE = "data/manifest.json"
FROZEN_FIXTURE_RESOURCE_DIRECTORY = ("data", "fixtures")
_DETERMINISTIC_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_CORE_PROPERTY_TIMESTAMP = b"1980-01-01T00:00:00Z"
_CORE_PROPERTY_TIMESTAMP_PATTERN = re.compile(
    rb"(<dcterms:(?:created|modified)\b[^>]*>)[^<]*(</dcterms:(?:created|modified)>)"
)


class FixtureRecord(TypedDict):
    id: str
    description: str
    sha256: str


class Corpus(TypedDict):
    kind: str
    not_ground_truth: bool
    recipe: str
    recipe_sha256: str
    fixtures: list[FixtureRecord]


class Manifest(TypedDict):
    schema_version: int
    corpus: Corpus
    lanes: dict[str, object]


# This is deliberately data, rather than a hash of Python source. It captures the
# frozen semantics of the corpus and survives formatting and implementation-only
# refactors. Changing it is a corpus change and requires explicit regeneration.
_FEATURE_CHECKS = [
    "slide-count",
    "slide-order-and-identifiers",
    "slide-relationship-graph",
    "shape-tree-semantics",
    "text-content-and-order",
    "text-run-formatting",
    "table-cell-grid-and-text",
    "table-cell-formatting",
    "bullet-paragraph-properties",
    "layout-geometry",
    "slide-layout-semantics",
    "notes-content",
    "notes-relationships",
    "media-part-presence",
    "media-exact-bytes",
    "theme-semantics",
    "master-semantics",
    "relationship-semantics",
    "opaque-package-part",
]
_DECLARED_LANES: dict[str, object] = {
    "feature-matrix": {"fixture": "mixed-60", "checks": _FEATURE_CHECKS},
    "template-mutation": {
        "fixture": "mixed-60",
        "edits": {
            "slide": 5,
            "table_shape": "PPTBenchTable",
            "cells": {"1,1": "UPDATED-TABLE-A", "2,2": "UPDATED-TABLE-B"},
            "bullet": {
                "slide": 7,
                "shape": "PPTBenchBullets",
                "paragraph": 1,
                "text": "UPDATED-BULLET",
            },
        },
    },
    "chart-data": {
        "fixture": "charts",
        "charts": [
            {
                "slide": 1,
                "kind": "category",
                "series": "Revenue",
                "categories": ["East", "42", "84"],
                "values": [42, 84, 126],
            },
            {
                "slide": 2,
                "kind": "xy",
                "series": "Trend",
                "points": [[5, 13], [21, 34], [55, 89]],
            },
            {
                "slide": 3,
                "kind": "bubble",
                "series": "Pipeline",
                "points": [[5, 13, 21], [34, 55, 89], [8, 13, 21]],
            },
        ],
    },
}


FROZEN_RECIPE: dict[str, object] = {
    "lanes": _DECLARED_LANES,
    "version": 2,
    "mixed-60": {
        "slides": 60,
        "title": "Synthetic mixed template slide {number:02d}",
        "table_slide": 5,
        "table_shape": "PPTBenchTable",
        "table_value": "S05R{row}C{column}",
        "bullet_slide": 7,
        "bullet_shape": "PPTBenchBullets",
        "bullets": ["Baseline bullet", "Second bullet", "Third bullet"],
        "title_style": {"bold": True, "size_pt": 16},
        "table_cell_margin_inches": 0.08,
        "bullet_levels": [0, 1, 1],
        "periodic_tables": [10, 20, 30, 40, 50, 60],
        "body_slides": [6, 12, 18, 24, 36, 42, 48, 54],
        "note": "Synthetic note retained separately from edit success.",
        "opaque_part": "ppt/unknown/pptbench.xml",
    },
    "charts": {
        "category": {
            "chart_type": "column-clustered",
            "series": "Revenue",
            "source": {
                "categories": ["North", "South", "West"],
                "values": [10, 20, 30],
            },
            "target": {"categories": ["East", "42", "84"], "values": [42, 84, 126]},
        },
        "xy": {
            "chart_type": "xy-scatter-lines",
            "series": "Trend",
            "source": [[1, 3], [2, 7], [3, 11]],
            "target": [[5, 13], [21, 34], [55, 89]],
        },
        "bubble": {
            "chart_type": "bubble",
            "series": "Pipeline",
            "source": [[1, 6, 4], [2, 9, 8], [3, 12, 12]],
            "target": [[5, 13, 21], [34, 55, 89], [8, 13, 21]],
        },
    },
}


def manifest() -> Manifest:
    """Load and validate the packaged, maintainer-controlled fixture manifest."""
    raw: object = json.loads(
        files("pptbench").joinpath(MANIFEST_RESOURCE).read_text(encoding="utf-8")
    )
    if not isinstance(raw, dict):
        raise TypeError("fixture manifest must be an object")
    schema_version = raw.get("schema_version")
    corpus = raw.get("corpus")
    lanes = raw.get("lanes")
    if (
        not isinstance(schema_version, int)
        or not isinstance(corpus, dict)
        or not isinstance(lanes, dict)
    ):
        raise TypeError("fixture manifest has invalid top-level fields")
    fixtures = corpus.get("fixtures")
    if not isinstance(fixtures, list):
        raise TypeError("fixture manifest corpus.fixtures must be a list")
    records: list[FixtureRecord] = []
    for item in fixtures:
        if not isinstance(item, dict):
            raise TypeError("fixture record must be an object")
        identifier = item.get("id")
        description = item.get("description")
        digest = item.get("sha256")
        if (
            not isinstance(identifier, str)
            or not isinstance(description, str)
            or not isinstance(digest, str)
        ):
            raise TypeError("fixture record has invalid fields")
        records.append({"id": identifier, "description": description, "sha256": digest})
    kind = corpus.get("kind")
    not_ground_truth = corpus.get("not_ground_truth")
    recipe = corpus.get("recipe")
    recipe_sha256 = corpus.get("recipe_sha256")
    if (
        not isinstance(kind, str)
        or not isinstance(not_ground_truth, bool)
        or not isinstance(recipe, str)
        or not isinstance(recipe_sha256, str)
    ):
        raise TypeError("fixture manifest corpus has invalid fields")
    return {
        "schema_version": schema_version,
        "corpus": {
            "kind": kind,
            "not_ground_truth": not_ground_truth,
            "recipe": recipe,
            "recipe_sha256": recipe_sha256,
            "fixtures": records,
        },
        "lanes": dict(lanes),
    }


def verify_frozen_data() -> None:
    data = manifest()
    if data["schema_version"] != 2:
        raise ValueError("unsupported fixture manifest schema")
    if data["lanes"] != _DECLARED_LANES:
        raise ValueError("fixture manifest lane contract does not match the frozen recipe")
    expected = data["corpus"]["recipe_sha256"]
    observed = _recipe_sha256()
    if expected != observed:
        raise ValueError(
            "fixture recipe digest does not match the frozen manifest; "
            "run regenerate_frozen_data() for an intentional corpus update"
        )
    _verify_frozen_fixture_resources(_fixture_hashes(data))


def regenerate_frozen_data(manifest_path: Path | None = None) -> str:
    """Explicitly regenerate frozen package resources and their manifest hashes.

    This is a maintainer operation. Normal materialization only copies the
    reviewed package resources, so a changed generator cannot affect a run.
    """
    destination = manifest_path or Path(__file__).with_name("data") / "manifest.json"
    raw: object = json.loads(destination.read_text(encoding="utf-8"))
    if not isinstance(raw, dict) or not isinstance(raw.get("corpus"), dict):
        raise TypeError("fixture manifest must contain a corpus object")
    corpus = raw["corpus"]
    fixtures = corpus.get("fixtures")
    if not isinstance(fixtures, list):
        raise TypeError("fixture manifest corpus.fixtures must be a list")
    with tempfile.TemporaryDirectory(prefix="pptbench-freeze-") as temporary:
        root = Path(temporary)
        generated_paths = {
            "mixed-60": root / "mixed-60.pptx",
            "charts": root / "charts.pptx",
        }
        _make_mixed_template(generated_paths["mixed-60"])
        _make_charts(generated_paths["charts"])
        generated = {
            identifier: generated_path.read_bytes()
            for identifier, generated_path in generated_paths.items()
        }
    seen: set[str] = set()
    for item in fixtures:
        if not isinstance(item, dict) or not isinstance(item.get("id"), str):
            raise TypeError("fixture record must contain an id")
        identifier = item["id"]
        payload = generated.get(identifier)
        if payload is None or identifier in seen:
            raise ValueError(f"unexpected or duplicate fixture id: {identifier!r}")
        item["sha256"] = hashlib.sha256(payload).hexdigest()
        _write_frozen_fixture(destination.parent, identifier, payload)
        seen.add(identifier)
    if seen != set(generated):
        raise ValueError("fixture manifest does not declare every generated fixture")
    corpus["recipe_sha256"] = _recipe_sha256()
    destination.write_text(json.dumps(raw, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return _recipe_sha256()


def fixture_root(run_root: Path) -> Path:
    return run_root / "fixtures"


def materialize(run_root: Path) -> dict[str, Path]:
    """Copy verified, immutable package fixture resources into a benchmark run."""
    verify_frozen_data()
    root = fixture_root(run_root)
    root.mkdir(parents=True, exist_ok=True)
    expected = _fixture_hashes(manifest())
    paths = {"mixed-60": root / "mixed-60.pptx", "charts": root / "charts.pptx"}
    for identifier, path in paths.items():
        if path.exists():
            if not path.is_file() or path.stat().st_size == 0:
                raise ValueError(f"fixture path is not a reusable nonempty file: {path}")
            if sha256_file(path) != expected[identifier]:
                raise ValueError(f"fixture hash mismatch; refusing to overwrite {path}")
            continue
        _materialize_new(path, _frozen_fixture_bytes(identifier), expected[identifier])
    return paths


def fixture_inventory(run_root: Path) -> list[dict[str, str]]:
    paths = materialize(run_root)
    expected = _fixture_hashes(manifest())
    return [
        {
            "id": identifier,
            "path": paths[identifier].name,
            "sha256": expected[identifier],
        }
        for identifier in ("mixed-60", "charts")
    ]


def verify_fixture(identifier: str, path: Path) -> None:
    """Require an input fixture to be the exact frozen artifact for its lane."""
    verify_frozen_data()
    expected = _fixture_hashes(manifest()).get(identifier)
    if expected is None:
        raise ValueError(f"unknown frozen fixture: {identifier}")
    if not path.is_file() or sha256_file(path) != expected:
        raise ValueError(f"fixture hash mismatch: {path}")


def _fixture_hashes(data: Manifest) -> dict[str, str]:
    hashes: dict[str, str] = {}
    for record in data["corpus"]["fixtures"]:
        digest = record["sha256"]
        if len(digest) != 64 or any(character not in "0123456789abcdef" for character in digest):
            raise ValueError(f"fixture {record['id']!r} has an invalid SHA-256 digest")
        hashes[record["id"]] = digest
    if set(hashes) != {"mixed-60", "charts"}:
        raise ValueError("fixture manifest must declare exactly mixed-60 and charts")
    return hashes


def _recipe_sha256() -> str:
    encoded = json.dumps(FROZEN_RECIPE, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _frozen_fixture_bytes(identifier: str) -> bytes:
    if identifier not in {"mixed-60", "charts"}:
        raise ValueError(f"unknown frozen fixture: {identifier}")
    resource = files("pptbench").joinpath(*FROZEN_FIXTURE_RESOURCE_DIRECTORY, f"{identifier}.pptx")
    if not resource.is_file():
        raise ValueError(f"missing frozen fixture resource: {identifier}")
    return resource.read_bytes()


def _verify_frozen_fixture_resources(expected: dict[str, str]) -> None:
    for identifier, digest in expected.items():
        observed = hashlib.sha256(_frozen_fixture_bytes(identifier)).hexdigest()
        if observed != digest:
            raise ValueError(f"frozen fixture resource hash mismatch: {identifier}")


def _write_frozen_fixture(data_root: Path, identifier: str, payload: bytes) -> None:
    destination = data_root / "fixtures" / f"{identifier}.pptx"
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{identifier}-", suffix=".pptx", dir=destination.parent
    )
    try:
        with os.fdopen(descriptor, "wb") as temporary:
            temporary.write(payload)
        Path(temporary_name).replace(destination)
    finally:
        Path(temporary_name).unlink(missing_ok=True)


def _materialize_new(path: Path, payload: bytes, expected_hash: str) -> None:
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{path.stem}-", suffix=".pptx", dir=path.parent
    )
    temporary = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "wb") as stream:
            stream.write(payload)
        if sha256_file(temporary) != expected_hash:
            raise ValueError("frozen fixture resource differs from the manifest hash")
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def _imports() -> tuple[Any, Any, Any, Any, Any, Any]:
    try:
        from pptx import Presentation
        from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.util import Inches
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise RuntimeError("python-pptx is required to materialize PPTBench fixtures") from exc
    return (
        Presentation,
        BubbleChartData,
        CategoryChartData,
        XyChartData,
        XL_CHART_TYPE,
        Inches,
    )


def _make_mixed_template(path: Path) -> None:
    Presentation, _, _, _, _, Inches = _imports()
    from pptx.util import Pt

    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    png = BytesIO(_PNG_1X1)
    for number in range(1, 61):
        slide = presentation.slides.add_slide(blank)
        title = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(8.5), Inches(0.4))
        title.name = f"PPTBenchTitle{number:02d}"
        title.text = f"Synthetic mixed template slide {number:02d}"
        title_run = title.text_frame.paragraphs[0].runs[0]
        title_run.font.bold = True
        title_run.font.size = Pt(16)
        if number == 5:
            table = slide.shapes.add_table(4, 4, Inches(0.6), Inches(1.1), Inches(8), Inches(2.5))
            table.name = "PPTBenchTable"
            for row in range(4):
                for column in range(4):
                    cell = table.table.cell(row, column)
                    cell.text = f"S05R{row}C{column}"
                    cell.margin_left = Inches(0.08)
        elif number == 7:
            body = slide.shapes.add_textbox(Inches(0.8), Inches(1.0), Inches(8), Inches(3))
            body.name = "PPTBenchBullets"
            first = body.text_frame.paragraphs[0]
            first.text = "Baseline bullet"
            first.level = 0
            second = body.text_frame.add_paragraph()
            second.text = "Second bullet"
            second.level = 1
            third = body.text_frame.add_paragraph()
            third.text = "Third bullet"
            third.level = 1
        elif number % 10 == 0:
            table = slide.shapes.add_table(3, 3, Inches(0.8), Inches(1), Inches(6), Inches(1.8))
            table.name = f"PPTBenchPeriodicTable{number:02d}"
            for row in range(3):
                for column in range(3):
                    table.table.cell(row, column).text = f"{number}/{row}/{column}"
        elif number % 6 == 0:
            body = slide.shapes.add_textbox(Inches(0.8), Inches(1), Inches(7.5), Inches(2))
            body.name = f"PPTBenchBody{number:02d}"
            body.text = f"Synthetic content for feature check {number}."
        if number == 2:
            slide.shapes.add_picture(png, Inches(8.4), Inches(0.4), Inches(0.3), Inches(0.3))
        if number == 1:
            slide.notes_slide.notes_text_frame.text = (
                "Synthetic note retained separately from edit success."
            )
    presentation.save(path)
    _normalise_pptx(
        path,
        {
            "ppt/unknown/pptbench.xml": (
                b'<pptbench:marker xmlns:pptbench="https://pptbench.local/unknown">opaque'
                b"</pptbench:marker>"
            )
        },
    )


def _make_charts(path: Path) -> None:
    (
        Presentation,
        BubbleChartData,
        CategoryChartData,
        XyChartData,
        XL_CHART_TYPE,
        Inches,
    ) = _imports()
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    category = CategoryChartData()
    category.categories = ["North", "South", "West"]
    category.add_series("Revenue", (10, 20, 30))
    slide = presentation.slides.add_slide(blank)
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(0.5),
        Inches(0.5),
        Inches(8),
        Inches(4),
        category,
    )
    xy = XyChartData()
    series = xy.add_series("Trend")
    for x, y in ((1, 3), (2, 7), (3, 11)):
        series.add_data_point(x, y)
    slide = presentation.slides.add_slide(blank)
    slide.shapes.add_chart(
        XL_CHART_TYPE.XY_SCATTER_LINES,
        Inches(0.5),
        Inches(0.5),
        Inches(8),
        Inches(4),
        xy,
    )
    bubble = BubbleChartData()
    series = bubble.add_series("Pipeline")
    for x, y, size in ((1, 6, 4), (2, 9, 8), (3, 12, 12)):
        series.add_data_point(x, y, size)
    slide = presentation.slides.add_slide(blank)
    slide.shapes.add_chart(
        XL_CHART_TYPE.BUBBLE, Inches(0.5), Inches(0.5), Inches(8), Inches(4), bubble
    )
    presentation.save(path)
    _normalise_pptx(path)


def _normalise_pptx(path: Path, additions: dict[str, bytes] | None = None) -> None:
    with zipfile.ZipFile(path) as source:
        parts = {
            item.filename: source.read(item.filename)
            for item in source.infolist()
            if not item.is_dir()
        }
    if additions:
        if set(parts).intersection(additions):
            raise ValueError("fixture addition collides with generated package part")
        parts.update(additions)
    _write_deterministic_zip(
        path,
        {name: _normalise_package_part(name, data) for name, data in parts.items()},
    )


def _normalise_package_part(name: str, data: bytes) -> bytes:
    if data.startswith(b"PK\x03\x04"):
        return _normalise_zip_bytes(data)
    if name.endswith((".xml", ".rels")):
        return _CORE_PROPERTY_TIMESTAMP_PATTERN.sub(
            rb"\g<1>" + _CORE_PROPERTY_TIMESTAMP + rb"\g<2>", data
        )
    return data


def _normalise_zip_bytes(data: bytes) -> bytes:
    with zipfile.ZipFile(BytesIO(data)) as source:
        parts = {
            item.filename: source.read(item.filename)
            for item in source.infolist()
            if not item.is_dir()
        }
    return _deterministic_zip_bytes(
        {name: _normalise_package_part(name, part) for name, part in parts.items()}
    )


def _write_deterministic_zip(path: Path, parts: dict[str, bytes]) -> None:
    staged = path.with_suffix(path.suffix + ".staged")
    try:
        staged.write_bytes(_deterministic_zip_bytes(parts))
        staged.replace(path)
    finally:
        staged.unlink(missing_ok=True)


def _deterministic_zip_bytes(parts: dict[str, bytes]) -> bytes:
    stream = BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for name, data in sorted(parts.items()):
            info = zipfile.ZipInfo(name, date_time=_DETERMINISTIC_ZIP_TIME)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 3
            info.external_attr = (stat.S_IFREG | 0o644) << 16
            archive.writestr(info, data, compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
    return stream.getvalue()


_PNG_1X1 = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000d49444154789c6360f8cfc0000004010100dffa8f4d0000000049454e44ae426082"
)
