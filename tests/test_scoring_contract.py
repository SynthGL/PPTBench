from __future__ import annotations

import hashlib
import os
import posixpath
import shutil
import subprocess
import sys
import zipfile
from collections.abc import Callable
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from pptx import Presentation
from pptx.chart.data import BubbleChartData, CategoryChartData, XyChartData

from pptbench import fixtures
from pptbench.fixtures import materialize, verify_frozen_data
from pptbench.scoring import score

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
_P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def _scored_failures(checks: list[dict[str, object]]) -> list[dict[str, object]]:
    return [check for check in checks if check["scored"] and check["outcome"] == "failure"]


def _replace_part(source: Path, output: Path, name: str, replacement: bytes) -> None:
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(output, "w") as changed:
        for item in original.infolist():
            changed.writestr(
                item,
                replacement if item.filename == name else original.read(item.filename),
            )


def _shape(root: ET.Element, name: str) -> ET.Element:
    for candidate in root.iter():
        properties = candidate.find(".//" + _P + "cNvPr")
        if properties is not None and properties.attrib.get("name") == name:
            return candidate
    raise AssertionError(f"missing shape {name}")


def _cell(root: ET.Element, row: int, column: int) -> ET.Element:
    table = _shape(root, "PPTBenchTable").find(".//" + _A + "tbl")
    assert table is not None
    return table.findall(_A + "tr")[row].findall(_A + "tc")[column]


def _set_text(element: ET.Element, value: str) -> None:
    nodes = list(element.iter(_A + "t"))
    assert len(nodes) == 1
    nodes[0].text = value


def _template_edit(source: Path, output: Path) -> None:
    with zipfile.ZipFile(source) as archive:
        parts = {item.filename: archive.read(item.filename) for item in archive.infolist()}
    slide5 = ET.fromstring(parts["ppt/slides/slide5.xml"])
    _set_text(_cell(slide5, 1, 1), "UPDATED-TABLE-A")
    _set_text(_cell(slide5, 2, 2), "UPDATED-TABLE-B")
    slide7 = ET.fromstring(parts["ppt/slides/slide7.xml"])
    paragraph = _shape(slide7, "PPTBenchBullets").findall(".//" + _A + "p")[1]
    _set_text(paragraph, "UPDATED-BULLET")
    parts["ppt/slides/slide5.xml"] = ET.tostring(slide5, encoding="utf-8")
    parts["ppt/slides/slide7.xml"] = ET.tostring(slide7, encoding="utf-8")
    with zipfile.ZipFile(output, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def _chart_edit(source: Path, output: Path) -> None:
    """Use python-pptx directly as a contestant producer, never as scorer oracle."""
    presentation = Presentation(source)
    category = CategoryChartData()
    category.categories = ["East", "42", "84"]
    category.add_series("Revenue", (42, 84, 126))
    presentation.slides[0].shapes[0].chart.replace_data(category)
    xy = XyChartData()
    xy_series = xy.add_series("Trend")
    for x, y in ((5, 13), (21, 34), (55, 89)):
        xy_series.add_data_point(x, y)
    presentation.slides[1].shapes[0].chart.replace_data(xy)
    bubble = BubbleChartData()
    bubble_series = bubble.add_series("Pipeline")
    for x, y, size in ((5, 13, 21), (34, 55, 89), (8, 13, 21)):
        bubble_series.add_data_point(x, y, size)
    presentation.slides[2].shapes[0].chart.replace_data(bubble)
    presentation.save(output)


def test_exact_template_targeting_and_collateral_semantics(tmp_path: Path) -> None:
    source = materialize(tmp_path)["mixed-60"]
    correct = tmp_path / "correct.pptx"
    _template_edit(source, correct)
    assert not _scored_failures(score("template-mutation", source, correct))

    swapped = tmp_path / "swapped.pptx"
    _template_edit(source, swapped)
    with zipfile.ZipFile(swapped) as archive:
        slide = ET.fromstring(archive.read("ppt/slides/slide5.xml"))
    _set_text(_cell(slide, 1, 1), "UPDATED-TABLE-B")
    _set_text(_cell(slide, 2, 2), "UPDATED-TABLE-A")
    _replace_part(
        swapped,
        swapped.with_suffix(".rewritten.pptx"),
        "ppt/slides/slide5.xml",
        ET.tostring(slide),
    )
    swapped.with_suffix(".rewritten.pptx").replace(swapped)
    assert _scored_failures(score("template-mutation", source, swapped))

    styled = tmp_path / "styled.pptx"
    _template_edit(source, styled)
    with zipfile.ZipFile(styled) as archive:
        slide = ET.fromstring(archive.read("ppt/slides/slide5.xml"))
    _shape(slide, "PPTBenchTable").set("rot", "60000")
    _replace_part(
        styled,
        styled.with_suffix(".rewritten.pptx"),
        "ppt/slides/slide5.xml",
        ET.tostring(slide),
    )
    styled.with_suffix(".rewritten.pptx").replace(styled)
    assert _scored_failures(score("template-mutation", source, styled))


def test_wrong_slide_injection_and_notes_mutation_fail(tmp_path: Path) -> None:
    source = materialize(tmp_path)["mixed-60"]
    injected = tmp_path / "injected.pptx"
    _template_edit(source, injected)
    with zipfile.ZipFile(injected) as archive:
        slide = ET.fromstring(archive.read("ppt/slides/slide1.xml"))
    ET.SubElement(slide, _A + "t").text = "UPDATED-TABLE-A"
    _replace_part(
        injected,
        injected.with_suffix(".rewritten.pptx"),
        "ppt/slides/slide1.xml",
        ET.tostring(slide),
    )
    injected.with_suffix(".rewritten.pptx").replace(injected)
    assert _scored_failures(score("template-mutation", source, injected))

    notes = tmp_path / "notes.pptx"
    _template_edit(source, notes)
    with zipfile.ZipFile(notes) as archive:
        note = ET.fromstring(archive.read("ppt/notesSlides/notesSlide1.xml"))
    next(element for element in note.iter(_A + "t")).text = "collateral note mutation"
    _replace_part(
        notes,
        notes.with_suffix(".rewritten.pptx"),
        "ppt/notesSlides/notesSlide1.xml",
        ET.tostring(note),
    )
    notes.with_suffix(".rewritten.pptx").replace(notes)
    assert _scored_failures(score("template-mutation", source, notes))


def test_chart_mapping_rejects_token_soup_wrong_series_and_wrong_workbook(
    tmp_path: Path,
) -> None:
    source = materialize(tmp_path)["charts"]
    correct = tmp_path / "charts-correct.pptx"
    _chart_edit(source, correct)
    assert not _scored_failures(score("chart-data", source, correct))

    token_soup = tmp_path / "token-soup.pptx"
    with zipfile.ZipFile(source) as archive:
        chart = ET.fromstring(archive.read("ppt/charts/chart1.xml"))
    for value in ("East", "42", "84", "126", "5", "13", "21", "34", "55", "89"):
        ET.SubElement(chart, _C + "v").text = value
    _replace_part(source, token_soup, "ppt/charts/chart1.xml", ET.tostring(chart))
    assert _scored_failures(score("chart-data", source, token_soup))

    wrong_series = tmp_path / "wrong-series.pptx"
    _chart_edit(source, wrong_series)
    with zipfile.ZipFile(wrong_series) as archive:
        chart = ET.fromstring(archive.read("ppt/charts/chart2.xml"))
    series_name = next(element for element in chart.iter(_C + "v") if element.text == "Trend")
    series_name.text = "Pipeline"
    _replace_part(
        wrong_series,
        wrong_series.with_suffix(".rewritten.pptx"),
        "ppt/charts/chart2.xml",
        ET.tostring(chart),
    )
    wrong_series.with_suffix(".rewritten.pptx").replace(wrong_series)
    assert _scored_failures(score("chart-data", source, wrong_series))

    wrong_workbook = tmp_path / "wrong-workbook.pptx"
    _chart_edit(source, wrong_workbook)
    with zipfile.ZipFile(wrong_workbook) as archive:
        workbook_name = next(
            name for name in archive.namelist() if name.startswith("ppt/embeddings/")
        )
        workbook = archive.read(workbook_name)
    with zipfile.ZipFile(BytesIO(workbook)) as original:
        workbook_parts = {name: original.read(name) for name in original.namelist()}
    worksheet_name = next(name for name in workbook_parts if name.startswith("xl/worksheets/"))
    worksheet = ET.fromstring(workbook_parts[worksheet_name])
    value = next(element for element in worksheet.iter() if element.tag.endswith("}v"))
    value.text = "999"
    workbook_parts[worksheet_name] = ET.tostring(worksheet, encoding="utf-8")
    rebuilt = BytesIO()
    with zipfile.ZipFile(rebuilt, "w") as archive:
        for name, data in workbook_parts.items():
            archive.writestr(name, data)
    _replace_part(
        wrong_workbook,
        wrong_workbook.with_suffix(".rewritten.pptx"),
        workbook_name,
        rebuilt.getvalue(),
    )
    wrong_workbook.with_suffix(".rewritten.pptx").replace(wrong_workbook)
    assert _scored_failures(score("chart-data", source, wrong_workbook))
    inline_string = tmp_path / "inline-string.pptx"
    _chart_edit(source, inline_string)
    with zipfile.ZipFile(inline_string) as archive:
        workbook_name = next(
            name for name in archive.namelist() if name.startswith("ppt/embeddings/")
        )
        workbook = archive.read(workbook_name)
    with zipfile.ZipFile(BytesIO(workbook)) as original:
        workbook_parts = {name: original.read(name) for name in original.namelist()}
    worksheet_name = next(name for name in workbook_parts if name.startswith("xl/worksheets/"))
    worksheet = ET.fromstring(workbook_parts[worksheet_name])
    east = next(cell for cell in worksheet.iter(_S + "c") if cell.attrib.get("r") == "A2")
    value = next(element for element in east if element.tag == _S + "v")
    east.remove(value)
    east.set("t", "inlineStr")
    inline = ET.SubElement(east, _S + "is")
    ET.SubElement(inline, _S + "t").text = "East"
    workbook_parts[worksheet_name] = ET.tostring(worksheet, encoding="utf-8")
    rebuilt = BytesIO()
    with zipfile.ZipFile(rebuilt, "w") as archive:
        for name, data in workbook_parts.items():
            archive.writestr(name, data)
    _replace_part(
        inline_string,
        inline_string.with_suffix(".rewritten.pptx"),
        workbook_name,
        rebuilt.getvalue(),
    )
    inline_string.with_suffix(".rewritten.pptx").replace(inline_string)
    assert not _scored_failures(score("chart-data", source, inline_string))

    style_corruption = tmp_path / "workbook-style-corruption.pptx"
    _chart_edit(source, style_corruption)
    with zipfile.ZipFile(style_corruption) as archive:
        workbook_name = next(
            name for name in archive.namelist() if name.startswith("ppt/embeddings/")
        )
        workbook = archive.read(workbook_name)
    with zipfile.ZipFile(BytesIO(workbook)) as original:
        workbook_parts = {name: original.read(name) for name in original.namelist()}
    worksheet_name = next(name for name in workbook_parts if name.startswith("xl/worksheets/"))
    worksheet = ET.fromstring(workbook_parts[worksheet_name])
    east = next(cell for cell in worksheet.iter(_S + "c") if cell.attrib.get("r") == "A2")
    east.set("s", "0")
    workbook_parts[worksheet_name] = ET.tostring(worksheet, encoding="utf-8")
    rebuilt = BytesIO()
    with zipfile.ZipFile(rebuilt, "w") as archive:
        for name, data in workbook_parts.items():
            archive.writestr(name, data)
    _replace_part(
        style_corruption,
        style_corruption.with_suffix(".rewritten.pptx"),
        workbook_name,
        rebuilt.getvalue(),
    )
    style_corruption.with_suffix(".rewritten.pptx").replace(style_corruption)
    assert any(
        check["name"] == "related-workbook-structure-and-formatting"
        and check["outcome"] == "failure"
        for check in score("chart-data", source, style_corruption)
    )


def test_reserialization_is_semantic_success_with_separate_byte_observation(
    tmp_path: Path,
) -> None:
    source = materialize(tmp_path)["mixed-60"]
    reserialized = tmp_path / "reserialized.pptx"
    with (
        zipfile.ZipFile(source) as original,
        zipfile.ZipFile(reserialized, "w", zipfile.ZIP_DEFLATED) as output,
    ):
        for name in reversed(original.namelist()):
            output.writestr(name, original.read(name))
    assert (
        hashlib.sha256(source.read_bytes()).digest()
        != hashlib.sha256(reserialized.read_bytes()).digest()
    )
    checks = score("feature-matrix", source, reserialized)
    assert not _scored_failures(checks)
    assert any(
        check["name"] == "raw-untouched-part-equality"
        and check["outcome"] == "success"
        and check["category"] == "byte-only"
        for check in checks
    )

    xml_reserialized = tmp_path / "xml-reserialized.pptx"
    with zipfile.ZipFile(source) as original:
        parts = {name: original.read(name) for name in original.namelist()}
    slide_name = "ppt/slides/slide1.xml"
    reserialized_slide = ET.tostring(
        ET.fromstring(parts[slide_name]), encoding="utf-8", xml_declaration=True
    )
    assert reserialized_slide != parts[slide_name]
    parts[slide_name] = reserialized_slide
    with zipfile.ZipFile(xml_reserialized, "w") as output:
        for name, data in parts.items():
            output.writestr(name, data)
    xml_checks = score("feature-matrix", source, xml_reserialized)
    assert not _scored_failures(xml_checks)
    assert any(
        check["name"] == "raw-untouched-part-equality"
        and check["outcome"] == "failure"
        and check["category"] == "byte-only"
        for check in xml_checks
    )


def test_invalid_zip_and_freeze_or_fixture_tampering_fail(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = materialize(tmp_path)["mixed-60"]
    invalid = tmp_path / "invalid.pptx"
    invalid.write_bytes(b"not a zip")
    assert _scored_failures(score("feature-matrix", source, invalid))

    monkeypatch.setattr(fixtures, "FROZEN_RECIPE", {"tampered": True})
    with pytest.raises(ValueError, match="digest"):
        verify_frozen_data()
    monkeypatch.undo()

    source.write_bytes(source.read_bytes() + b"tampered")
    with pytest.raises(ValueError, match="hash mismatch"):
        materialize(tmp_path)


def _zip_bytes(parts: dict[str, bytes]) -> bytes:
    output = BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    return output.getvalue()


def _timestamped_core_properties(timestamp: str) -> bytes:
    return (
        '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
        'xmlns:dcterms="http://purl.org/dc/terms/" '
        'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
        f'<dcterms:created xsi:type="dcterms:W3CDTF">{timestamp}</dcterms:created>'
        f'<dcterms:modified xsi:type="dcterms:W3CDTF">{timestamp}</dcterms:modified>'
        "</cp:coreProperties>"
    ).encode()


def _timestamped_package(timestamp: str) -> bytes:
    workbook = _zip_bytes(
        {
            "docProps/core.xml": _timestamped_core_properties(timestamp),
            "xl/workbook.xml": b"<workbook/>",
        }
    )
    return _zip_bytes(
        {
            "docProps/core.xml": _timestamped_core_properties(timestamp),
            "ppt/embeddings/workbook.xlsx": workbook,
        }
    )


def _regenerated_fixture_bytes(destination: Path) -> dict[str, bytes]:
    source_root = Path(__file__).parents[1] / "src"
    manifest_path = destination / "manifest.json"
    destination.mkdir()
    shutil.copyfile(source_root / "pptbench" / "data" / "manifest.json", manifest_path)
    environment = os.environ | {
        "PYTHONPATH": os.pathsep.join(
            part for part in (str(source_root), os.environ.get("PYTHONPATH")) if part
        )
    }
    subprocess.run(
        [
            sys.executable,
            "-c",
            (
                "from pathlib import Path; "
                "from pptbench.fixtures import regenerate_frozen_data; "
                "import sys; "
                "regenerate_frozen_data(Path(sys.argv[1]))"
            ),
            str(manifest_path),
        ],
        check=True,
        cwd=source_root.parent,
        env=environment,
        capture_output=True,
        text=True,
    )
    return {
        identifier: (destination / "fixtures" / f"{identifier}.pptx").read_bytes()
        for identifier in ("mixed-60", "charts")
    }


def _assert_deterministic_archives(value: bytes) -> None:
    with zipfile.ZipFile(BytesIO(value)) as archive:
        for info in archive.infolist():
            assert info.date_time == fixtures._DETERMINISTIC_ZIP_TIME
            assert info.create_system == 3
            assert info.external_attr == (0o100644 << 16)
            part = archive.read(info.filename)
            if info.filename == "docProps/core.xml":
                assert part.count(fixtures._CORE_PROPERTY_TIMESTAMP) == 2
            if part.startswith(b"PK\x03\x04"):
                _assert_deterministic_archives(part)


_LANE_FIXTURES = {
    "feature-matrix": "mixed-60",
    "template-mutation": "mixed-60",
    "chart-data": "charts",
}


def test_frozen_fixture_resources_are_reproducible_across_timestamps_and_processes(
    tmp_path: Path,
) -> None:
    assert fixtures._normalise_zip_bytes(
        _timestamped_package("2001-02-03T04:05:06Z")
    ) == fixtures._normalise_zip_bytes(_timestamped_package("2030-04-05T06:07:08Z"))

    first = _regenerated_fixture_bytes(tmp_path / "first")
    second = _regenerated_fixture_bytes(tmp_path / "second")
    materialized = materialize(tmp_path / "run")
    for identifier in ("mixed-60", "charts"):
        frozen = fixtures._frozen_fixture_bytes(identifier)
        assert first[identifier] == second[identifier] == frozen
        assert materialized[identifier].read_bytes() == frozen
        _assert_deterministic_archives(first[identifier])


# Equivalent rewrites ------------------------------------------------------------------
#
# A writer may re-serialize a package into an equivalent form. These helpers model the
# most aggressive such writer: every part renamed and moved, every relationship id
# renumbered (and relationship order reversed), every XML part re-encoded with new
# namespace prefixes and schema defaults written out, ZIP order reversed, and save-time
# metadata regenerated, recursively through embedded packages. Scored checks must not
# see any of it, yet must still catch a real content loss made before the rewrite.

_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_CT = "{http://schemas.openxmlformats.org/package/2006/content-types}"
_CORE = "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}"
_EXTENDED = "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}"
_DCTERMS = "{http://purl.org/dc/terms/}"


def _rels_name(source: str) -> str:
    directory, base = posixpath.split(source)
    return posixpath.join(directory, "_rels", base + ".rels")


def _is_xml(name: str, data: bytes) -> bool:
    return name.endswith((".xml", ".rels")) or data.lstrip().startswith(b"<?xml")


def _regenerate_save_time_metadata(root: ET.Element) -> None:
    if root.tag == _CORE + "coreProperties":
        for element in root:
            if element.tag == _DCTERMS + "modified":
                element.text = "2031-07-08T09:10:11Z"
            elif element.tag == _CORE + "revision":
                element.text = "42"
        ET.SubElement(root, _CORE + "lastModifiedBy").text = "Equivalent Rewriter"
    elif root.tag == _EXTENDED + "Properties":
        for element in root:
            if element.tag in {_EXTENDED + "AppVersion", _EXTENDED + "TotalTime"}:
                element.text = "99.0000"


def _write_schema_defaults(root: ET.Element) -> None:
    for element in root.iter():
        if element.tag == _P + "cNvSpPr":
            element.set("txBox", {"1": "true"}.get(element.get("txBox", ""), "false"))
        elif element.tag == _A + "tcPr":
            element.attrib.setdefault("marL", "91440")
        elif element.tag == _P + "cNvPr":
            element.attrib.setdefault("hidden", "false")


def _equivalent_rewrite(parts: dict[str, bytes]) -> dict[str, bytes]:
    plumbing = {"[Content_Types].xml"}
    content = [
        name for name in sorted(parts) if name not in plumbing and not name.endswith(".rels")
    ]
    renamed = {
        name: f"moved/{index:03d}-{hashlib.sha256(name.encode()).hexdigest()[:6]}"
        + posixpath.splitext(name)[1]
        for index, name in enumerate(reversed(content))
    }
    result: dict[str, bytes] = {}
    id_maps: dict[str, dict[str, str]] = {}
    for source in ["", *content]:
        rels_name = _rels_name(source) if source else "_rels/.rels"
        if rels_name not in parts:
            continue
        relationships = ET.fromstring(parts[rels_name])
        entries = list(relationships)
        id_map: dict[str, str] = {}
        for element in entries:
            relationships.remove(element)
        for number, element in enumerate(reversed(entries)):
            id_map[element.attrib["Id"]] = f"renumbered{number}"
            element.set("Id", id_map[element.attrib["Id"]])
            if element.get("TargetMode") != "External":
                target = posixpath.normpath(
                    posixpath.join(posixpath.dirname(source), element.attrib["Target"])
                    if not element.attrib["Target"].startswith("/")
                    else element.attrib["Target"].lstrip("/")
                )
                new_source = renamed.get(source, "")
                new_target = renamed.get(target, target)
                element.set(
                    "Target",
                    posixpath.relpath(new_target, posixpath.dirname(new_source))
                    if new_source
                    else "/" + new_target,
                )
            relationships.append(element)
        id_maps[source] = id_map
        new_rels = _rels_name(renamed[source]) if source else "_rels/.rels"
        result[new_rels] = ET.tostring(relationships, encoding="utf-8")
    for name in content:
        data = parts[name]
        if data.startswith(b"PK\x03\x04"):
            data = _zip_bytes(_equivalent_rewrite(_unzip(data)))
        elif _is_xml(name, data):
            root = ET.fromstring(data)
            id_map = id_maps.get(name, {})
            for element in root.iter():
                for attribute, value in element.attrib.items():
                    if attribute.startswith(_R) and value in id_map:
                        element.set(attribute, id_map[value])
            _regenerate_save_time_metadata(root)
            _write_schema_defaults(root)
            data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        result[renamed[name]] = data
    content_types = ET.fromstring(parts["[Content_Types].xml"])
    for element in content_types:
        if element.tag == _CT + "Override":
            old = element.attrib["PartName"].lstrip("/")
            element.set("PartName", "/" + renamed.get(old, old))
    result["[Content_Types].xml"] = ET.tostring(content_types, encoding="utf-8")
    return dict(reversed(list(result.items())))


def _unzip(value: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(BytesIO(value)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def _edit_xml(parts: dict[str, bytes], name: str, edit: Callable[[ET.Element], None]) -> None:
    root = ET.fromstring(parts[name])
    edit(root)
    parts[name] = ET.tostring(root, encoding="utf-8")


def _chart_named(parts: dict[str, bytes], kind: str) -> str:
    return next(
        name
        for name, data in parts.items()
        if name.startswith("ppt/charts/chart") and kind.encode() in data
    )


def _chart_workbook(parts: dict[str, bytes], chart: str) -> str:
    relationships = ET.fromstring(parts[_rels_name(chart)])
    target = next(
        element.attrib["Target"]
        for element in relationships
        if element.attrib["Type"].endswith("/package")
    )
    return posixpath.normpath(posixpath.join(posixpath.dirname(chart), target))


def _edit_workbook(
    parts: dict[str, bytes],
    chart_kind: str,
    prefix: str,
    edit: Callable[[ET.Element], None],
) -> None:
    workbook = _chart_workbook(parts, _chart_named(parts, chart_kind))
    inner = _unzip(parts[workbook])
    _edit_xml(inner, next(name for name in inner if name.startswith(prefix)), edit)
    parts[workbook] = _zip_bytes(inner)


def _produce(lane: str, source: Path, tmp_path: Path) -> dict[str, bytes]:
    output = tmp_path / "produced.pptx"
    if lane == "template-mutation":
        _template_edit(source, output)
    elif lane == "chart-data":
        _chart_edit(source, output)
    else:
        shutil.copyfile(source, output)
    return _unzip(output.read_bytes())


def _score_parts(
    lane: str, source: Path, parts: dict[str, bytes], tmp_path: Path
) -> list[dict[str, object]]:
    output = tmp_path / "rewritten.pptx"
    output.write_bytes(_zip_bytes(parts))
    return score(lane, source, output)


def _set_first_text(root: ET.Element, value: str) -> None:
    next(element for element in root.iter(_A + "t")).text = value


def _add_external_link(root: ET.Element) -> None:
    ET.SubElement(
        root,
        _REL + "Relationship",
        {
            "Id": "rIdAdded",
            "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            "Target": "https://example.invalid/",
            "TargetMode": "External",
        },
    )


def _retarget(rels: str, suffix: str, target: str) -> Callable[[dict[str, bytes]], None]:
    def edit(parts: dict[str, bytes]) -> None:
        def change(root: ET.Element) -> None:
            next(e for e in root if e.attrib["Type"].endswith(suffix)).set("Target", target)

        _edit_xml(parts, rels, change)

    return edit


def _drop_image_relationship(parts: dict[str, bytes]) -> None:
    def drop(root: ET.Element) -> None:
        root.remove(next(e for e in root if e.attrib["Type"].endswith("/image")))

    _edit_xml(parts, "ppt/slides/_rels/slide2.xml.rels", drop)


def _change_image(parts: dict[str, bytes]) -> None:
    parts["ppt/media/image1.png"] += b"\x00re-encoded"


def _drop_opaque(parts: dict[str, bytes]) -> None:
    del parts["ppt/unknown/pptbench.xml"]


def _change_opaque(parts: dict[str, bytes]) -> None:
    _edit_xml(parts, "ppt/unknown/pptbench.xml", lambda root: setattr(root, "text", "lost"))


def _change_notes(parts: dict[str, bytes]) -> None:
    _edit_xml(parts, "ppt/notesSlides/notesSlide1.xml", lambda r: _set_first_text(r, "lost"))


def _change_slide_text(parts: dict[str, bytes]) -> None:
    _edit_xml(parts, "ppt/slides/slide1.xml", lambda r: _set_first_text(r, "lost"))


def _change_theme(parts: dict[str, bytes]) -> None:
    def recolor(root: ET.Element) -> None:
        next(root.iter(_A + "srgbClr")).set("val", "FF00FF")

    _edit_xml(parts, "ppt/theme/theme1.xml", recolor)


def _change_master(parts: dict[str, bytes]) -> None:
    def hide(root: ET.Element) -> None:
        next(root.iter(_P + "cNvPr")).set("name", "lost")

    _edit_xml(parts, "ppt/slideMasters/slideMaster1.xml", hide)


def _change_layout(parts: dict[str, bytes]) -> None:
    def rename(root: ET.Element) -> None:
        next(root.iter(_P + "cSld")).set("name", "lost")

    _edit_xml(parts, "ppt/slideLayouts/slideLayout1.xml", rename)


def _link_notes(parts: dict[str, bytes]) -> None:
    _edit_xml(parts, "ppt/notesSlides/_rels/notesSlide1.xml.rels", _add_external_link)


def _restyle_bar_chart(parts: dict[str, bytes]) -> None:
    def flip(root: ET.Element) -> None:
        next(root.iter(_C + "barDir")).set("val", "bar")

    _edit_xml(parts, _chart_named(parts, "barChart"), flip)


def _change_series_header(parts: dict[str, bytes]) -> None:
    def rename(root: ET.Element) -> None:
        next(t for t in root.iter(_S + "t") if t.text == "Revenue").text = "Expenses"

    _edit_workbook(parts, "barChart", "xl/sharedStrings", rename)


def _restyle_workbook_cell(parts: dict[str, bytes]) -> None:
    def restyle(root: ET.Element) -> None:
        next(c for c in root.iter(_S + "c") if c.get("r") == "A2").set("s", "0")

    _edit_workbook(parts, "barChart", "xl/worksheets/", restyle)


def _share_workbook(parts: dict[str, bytes]) -> None:
    shared = _chart_workbook(parts, _chart_named(parts, "barChart"))
    scatter = _chart_named(parts, "scatterChart")
    _retarget(_rels_name(scatter), "/package", "/" + shared)(parts)


_LOSSES: list[tuple[str, str, Callable[[dict[str, bytes]], None]]] = [
    ("feature-matrix", "media-part-presence", _drop_image_relationship),
    ("feature-matrix", "media-exact-bytes", _change_image),
    ("feature-matrix", "theme-semantics", _change_theme),
    ("feature-matrix", "master-semantics", _change_master),
    ("feature-matrix", "slide-layout-semantics", _change_layout),
    ("feature-matrix", "notes-content", _change_notes),
    ("feature-matrix", "notes-relationships", _link_notes),
    ("feature-matrix", "relationship-semantics", _link_notes),
    (
        "feature-matrix",
        "slide-relationship-graph",
        _retarget(
            "ppt/slides/_rels/slide1.xml.rels",
            "/slideLayout",
            "../slideLayouts/slideLayout1.xml",
        ),
    ),
    ("feature-matrix", "opaque-package-part", _change_opaque),
    ("template-mutation", "only-declared-text-nodes-changed", _change_slide_text),
    ("template-mutation", "notes-preserved", _change_notes),
    ("template-mutation", "media-preserved", _change_image),
    ("template-mutation", "relationships-preserved", _link_notes),
    ("template-mutation", "opaque-parts-preserved", _drop_opaque),
    ("chart-data", "chart-topology-and-relations", _share_workbook),
    ("chart-data", "related-workbook-values", _change_series_header),
    ("chart-data", "related-workbook-structure-and-formatting", _restyle_workbook_cell),
    ("chart-data", "chart-formatting-preserved", _restyle_bar_chart),
]


@pytest.mark.parametrize("lane", ["feature-matrix", "template-mutation", "chart-data"])
def test_equivalent_rewrite_scores_like_the_original(lane: str, tmp_path: Path) -> None:
    source = materialize(tmp_path)[_LANE_FIXTURES[lane]]
    produced = _produce(lane, source, tmp_path)
    rewritten = _equivalent_rewrite(produced)
    assert not set(rewritten) & set(produced) - {"[Content_Types].xml", "_rels/.rels"}
    checks = _score_parts(lane, source, rewritten, tmp_path)
    assert not _scored_failures(checks)
    assert {check["name"] for check in checks if check["scored"]} >= {
        name for check_lane, name, _ in _LOSSES if check_lane == lane
    }
    assert all(
        check["outcome"] == "failure"
        for check in checks
        if check["name"] == "raw-untouched-part-equality"
    )


@pytest.mark.parametrize(
    ("lane", "check", "lose"),
    _LOSSES,
    ids=[f"{lane}:{check}" for lane, check, _ in _LOSSES],
)
def test_content_loss_survives_an_equivalent_rewrite(
    lane: str,
    check: str,
    lose: Callable[[dict[str, bytes]], None],
    tmp_path: Path,
) -> None:
    source = materialize(tmp_path)[_LANE_FIXTURES[lane]]
    produced = _produce(lane, source, tmp_path)
    lose(produced)
    checks = _score_parts(lane, source, _equivalent_rewrite(produced), tmp_path)
    assert check in {failure["name"] for failure in _scored_failures(checks)}
